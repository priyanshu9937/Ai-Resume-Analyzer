from dataclasses import dataclass
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from pypdf import PdfReader


class DocumentParseError(ValueError):
    pass


MAX_DOCX_UNCOMPRESSED_BYTES = 20 * 1024 * 1024
MAX_DOCX_PARTS = 1000
MAX_PDF_PAGES = 100


@dataclass(frozen=True)
class ExtractedDocument:
    text: str
    page_count: int | None = None
    paragraph_count: int | None = None


def extract_text(content: bytes, extension: str) -> ExtractedDocument:
    try:
        if extension == ".pdf":
            reader = PdfReader(BytesIO(content))
            page_count = len(reader.pages)
            if page_count > MAX_PDF_PAGES:
                raise DocumentParseError("The PDF contains too many pages")
            return ExtractedDocument("\n".join(page.extract_text() or "" for page in reader.pages), page_count=page_count)
        if extension == ".docx":
            with ZipFile(BytesIO(content)) as archive:
                entries = archive.infolist()
                if len(entries) > MAX_DOCX_PARTS or sum(entry.file_size for entry in entries) > MAX_DOCX_UNCOMPRESSED_BYTES:
                    raise DocumentParseError("The DOCX archive exceeds the parsing limits")
            document = Document(BytesIO(content))
            return ExtractedDocument("\n".join(paragraph.text for paragraph in document.paragraphs), paragraph_count=len(document.paragraphs))
    except DocumentParseError:
        raise
    except Exception as exc:
        raise DocumentParseError("The uploaded document could not be parsed") from exc
    raise DocumentParseError("Unsupported document type")
