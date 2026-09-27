import asyncio
from io import BytesIO
from threading import Event

import pytest
from docx import Document
from httpx import ASGITransport, AsyncClient
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.core.dependencies import get_llm_provider
from app.core.config import Settings
from app.db.session import init_db
from app.main import app
from app.services.llm import parse_llm_response
from app.services.upload import parse_upload


def make_docx() -> bytes:
    document = Document()
    document.add_heading("Resume", level=1)
    document.add_paragraph("Python FastAPI engineer with SQL experience")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def make_pdf() -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = writer._add_object(DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")}))
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
    content = DecodedStreamObject()
    content.set_data(b"BT /F1 12 Tf 72 720 Td (Python FastAPI resume) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(content)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


@pytest.fixture(autouse=True)
async def initialize_database() -> None:
    await init_db()


@pytest.mark.asyncio
async def test_health_and_docs() -> None:
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        assert (await client.get("/health")).status_code == 200
        assert (await client.get("/docs")).status_code == 200


@pytest.mark.asyncio
async def test_frontend_and_assets_are_served() -> None:
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        page = await client.get("/")
        styles = await client.get("/styles.css")
        config = await client.get("/config.js")
        script = await client.get("/app.js")
    assert page.status_code == 200
    assert "Sift | Resume review" in page.text
    assert page.headers["x-content-type-options"] == "nosniff"
    assert page.headers["x-frame-options"] == "DENY"
    assert styles.status_code == 200
    assert "--forest: #17251f" in styles.text
    assert config.status_code == 200
    assert 'window.SIFT_API_BASE_URL = ""' in config.text
    assert script.status_code == 200
    assert "function renderAnalysis" in script.text
    assert "apiBaseUrl" in script.text


def test_production_requires_strong_api_token() -> None:
    with pytest.raises(ValueError, match="API_ACCESS_TOKEN must be configured"):
        Settings(environment="production")
    with pytest.raises(ValueError, match="at least 32 characters"):
        Settings(api_access_token="short")


def test_postgres_urls_use_async_driver() -> None:
    assert Settings(database_url="postgres://user:pass@host/db").database_url == "postgresql+asyncpg://user:pass@host/db"
    assert Settings(database_url="postgresql://user:pass@host/db").database_url == "postgresql+asyncpg://user:pass@host/db"


@pytest.mark.asyncio
async def test_api_requires_configured_access_token(monkeypatch: pytest.MonkeyPatch) -> None:
    token = "local-test-token-with-at-least-32-characters"
    monkeypatch.setattr("app.main.settings.api_access_token", token)
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        unauthorized = await client.post("/api/v1/analysis/missing")
        authorized = await client.post("/api/v1/analysis/missing", headers={"Authorization": f"Bearer {token}"})
    assert unauthorized.status_code == 401
    assert unauthorized.json()["error"]["code"] == "unauthorized"
    assert authorized.status_code == 404


@pytest.mark.asyncio
async def test_upload_rejects_wrong_mime() -> None:
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        response = await client.post("/api/v1/resumes/upload", files={"file": ("resume.pdf", b"%PDF test", "text/plain")})
    assert response.status_code == 415
    assert response.json()["error"]["code"] == "unsupported_media_type"


@pytest.mark.asyncio
async def test_upload_accepts_pdf_and_docx_with_metadata() -> None:
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        pdf = await client.post("/api/v1/resumes/upload", files={"file": ("resume.pdf", make_pdf(), "application/pdf")})
        docx = await client.post("/api/v1/resumes/upload", files={"file": ("resume.docx", make_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert pdf.status_code == 201
    assert pdf.json()["metadata"]["page_count"] == 1
    assert docx.status_code == 201
    assert docx.json()["metadata"]["paragraph_count"] >= 2


@pytest.mark.asyncio
async def test_slow_resume_parsing_does_not_block_other_requests(monkeypatch: pytest.MonkeyPatch) -> None:
    parser_started = Event()
    allow_parser_to_finish = Event()

    def slow_parser(content: bytes, extension: str):
        parser_started.set()
        if not allow_parser_to_finish.wait(timeout=5):
            raise TimeoutError("Test parser was not released")
        return parse_upload(content, extension)

    monkeypatch.setattr("app.api.routes.resumes.parse_upload", slow_parser)
    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            upload_task = asyncio.create_task(client.post(
                "/api/v1/resumes/upload",
                files={"file": ("resume.docx", make_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            ))
            assert await asyncio.to_thread(parser_started.wait, 2)
            health = await client.get("/health")
            allow_parser_to_finish.set()
            upload = await upload_task
        assert health.status_code == 200
        assert upload.status_code == 201
    finally:
        allow_parser_to_finish.set()


@pytest.mark.asyncio
async def test_upload_rejects_empty_unsupported_and_oversized_files() -> None:
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        empty = await client.post("/api/v1/resumes/upload", files={"file": ("resume.pdf", b"", "application/pdf")})
        unsupported = await client.post("/api/v1/resumes/upload", files={"file": ("resume.txt", b"resume", "text/plain")})
        oversized = await client.post("/api/v1/resumes/upload", files={"file": ("resume.pdf", b"%PDF" + b"x" * (5 * 1024 * 1024), "application/pdf")})
    assert empty.status_code == 422
    assert unsupported.status_code == 415
    assert oversized.status_code == 413


@pytest.mark.asyncio
async def test_analysis_uses_stored_result_without_second_llm_call() -> None:
    class CountingProvider:
        calls = 0

        async def analyze_resume(self, text: str):
            self.calls += 1
            return parse_llm_response('{"skills": ["python"]}', text)

    provider = CountingProvider()
    app.dependency_overrides[get_llm_provider] = lambda: provider
    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            upload = await client.post("/api/v1/resumes/upload", files={"file": ("resume.docx", make_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
            resume_id = upload.json()["resume"]["id"]
            first = await client.post(f"/api/v1/analysis/{resume_id}")
            second = await client.post(f"/api/v1/analysis/{resume_id}")
        assert first.status_code == 200
        assert second.json()["source"] == "stored"
        assert provider.calls == 1
        analysis = first.json()["analysis"]
        assert any("python" in strength.lower() for strength in analysis["strengths"])
        assert any("Experience section" in improvement for improvement in analysis["improvements"])
    finally:
        app.dependency_overrides.clear()


@pytest.mark.asyncio
async def test_job_match_validates_and_persists_result() -> None:
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        upload = await client.post("/api/v1/resumes/upload", files={"file": ("resume.docx", make_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
        resume_id = upload.json()["resume"]["id"]
        invalid = await client.post("/api/v1/job-match", json={"resume_id": resume_id, "job_description": "too short"})
        match = await client.post("/api/v1/job-match", json={"resume_id": resume_id, "job_description": "Python FastAPI engineer with SQL experience"})
    assert invalid.status_code == 422
    assert match.status_code == 200
    assert 0 <= match.json()["match_score"] <= 100


@pytest.mark.asyncio
async def test_get_analysis_does_not_generate_missing_analysis() -> None:
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        upload = await client.post("/api/v1/resumes/upload", files={"file": ("resume.docx", make_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
        resume_id = upload.json()["resume"]["id"]
        response = await client.get(f"/api/v1/analysis/{resume_id}")
    assert response.status_code == 404
    assert response.json()["error"]["code"] == "analysis_not_found"


@pytest.mark.asyncio
async def test_missing_resources_return_404() -> None:
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        analysis = await client.post("/api/v1/analysis/missing")
        match = await client.post("/api/v1/job-match", json={"resume_id": "missing", "job_description": "Python backend engineer with APIs"})
    assert analysis.status_code == 404
    assert match.status_code == 404
